import os
from typing import Optional
from docx import Document as DocxDocument
import PyPDF2


def detect_file_type(filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".txt":
        return "txt"
    if ext == ".docx":
        return "docx"
    if ext == ".pdf":
        return "pdf"

    return "unknown"


def extract_text_from_txt(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def extract_text_from_docx(file_path: str) -> str:
    doc = DocxDocument(file_path)
    parts = []

    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                cell_text = " ".join(
                    p.text.strip() for p in cell.paragraphs if p.text.strip()
                ).strip()
                cells.append(cell_text)
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def extract_text_from_pdf(file_path: str) -> str:
    text = []
    with open(file_path, "rb") as f:
        reader = PyPDF2.PdfReader(f)
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text.append(page_text)
    return "\n".join(text)


def extract_text_from_file(file_path: str, file_type: str) -> str:
    if file_type == "txt":
        return extract_text_from_txt(file_path)
    if file_type == "docx":
        return extract_text_from_docx(file_path)
    if file_type == "pdf":
        return extract_text_from_pdf(file_path)
    return ""


def open_docx(file_path: str) -> Optional[DocxDocument]:
    if not file_path.lower().endswith(".docx"):
        return None
    return DocxDocument(file_path)